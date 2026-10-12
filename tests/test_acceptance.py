import io
import json
import re
from pathlib import Path

import openpyxl
import pandas as pd
import pytest
from flask import Flask
from sqlalchemy import event

from app.app import create_app
from app.config import DevelopmentConfig, ProductionConfig, TestingConfig
from app.models_db import AuditLog, Dataset, PredictionResult, User, db
from app.utils.maintenance_rules import get_maintenance_recommendation
from app.utils.report_generator import generate_excel, generate_pdf
from app.utils.translations import SUPPORTED_LANGS, TRANSLATIONS, t
from ml.predict import build_input_row, predict_single
from ml.preprocessing import (
    CATEGORICAL_LEVELS,
    NUMERIC_FEATURES,
    clean_for_inference,
    clean_for_upload,
    is_supported_string_identifier,
    normalize_columns,
)


@pytest.fixture
def app():
    application = create_app('testing')
    application.config.update(TESTING=True, WTF_CSRF_ENABLED=True)
    with application.app_context():
        db.drop_all()
        db.create_all()
    yield application
    with application.app_context():
        db.session.remove()
        db.drop_all()


@pytest.fixture
def client(app):
    return app.test_client()


def _csrf(client, path):
    response = client.get(path)
    assert response.status_code == 200
    token = re.search(
        rb'name="csrf_token" value="([^"]+)"', response.data
    )
    assert token is not None
    return token.group(1).decode()


def _valid_row(**overrides):
    values = {
        'temperature': 25,
        'irradiance': 500,
        'humidity': 50,
        'panel_age': 10,
        'maintenance_count': 4,
        'soiling_ratio': 0.8,
        'voltage': 12,
        'current': 1.5,
        'module_temperature': 30,
        'cloud_coverage': 25,
        'wind_speed': 7,
        'pressure': 1013,
        'string_id': 'A1',
        'error_code': 'E00',
        'installation_type': 'fixed',
    }
    values.update(overrides)
    return values


def _register(client, username, password='Solar@2026x',
              email=None):
    token = _csrf(client, '/register')
    return client.post('/register', data={
        'csrf_token': token,
        'username': username,
        'email': email or f'{username}@example.test',
        'password': password,
        'password_confirm': password,
    }, follow_redirects=True)


def _authenticate_client(app, client, username='owner'):
    user_id = _create_user(app, username)
    with client.session_transaction() as session:
        session['_user_id'] = str(user_id)
        session['_fresh'] = True
    return user_id


def _create_user(app, username):
    with app.app_context():
        user = User(username=username, email=f'{username}@example.test')
        user.set_password('Solar@2026x')
        db.session.add(user)
        db.session.commit()
        return user.id


def _create_prediction(app, user_id, efficiency=0.45, soiling_ratio=0.5):
    with app.app_context():
        result = PredictionResult(
            user_id=user_id,
            input_data=json.dumps(_valid_row(soiling_ratio=soiling_ratio)),
            predicted_efficiency=efficiency,
            maintenance_recommendation=(
                'Schedule panel cleaning — heavy soiling detected'
            ),
        )
        db.session.add(result)
        db.session.commit()
        return result.id


def test_fixed_categories_change_predictions():
    baseline = _valid_row(
        string_id='A1', error_code='E00', installation_type='fixed'
    )
    alternate = _valid_row(
        string_id='B2', error_code='E01', installation_type='tracking'
    )
    assert predict_single(baseline) != pytest.approx(predict_single(alternate))


def test_single_category_csv_uses_fixed_levels():
    csv_text = pd.DataFrame([_valid_row(string_id='B2')]).to_csv(index=False)
    raw = pd.read_csv(io.StringIO(csv_text))
    encoded = clean_for_inference(raw)
    assert encoded.iloc[0]['string_id_B2'] == 1
    assert all(
        encoded.iloc[0][f'string_id_{level}'] == 0
        for level in CATEGORICAL_LEVELS['string_id'][1:]
        if level != 'B2'
    )
    row = build_input_row(_valid_row(string_id='B2'), list(encoded.columns))
    assert row.iloc[0]['string_id_B2'] == 1


def test_unknown_categories_raise_a_clear_validation_error():
    with pytest.raises(ValueError, match='Unknown string_id value'):
        clean_for_inference(pd.DataFrame([_valid_row(string_id='ZZ')]))


def _post_csv(client, csv_data, filename='panels.csv', follow_redirects=False):
    token = _csrf(client, '/upload')
    return client.post('/upload', data={
        'csrf_token': token,
        'file': (io.BytesIO(csv_data), filename),
    }, content_type='multipart/form-data', follow_redirects=follow_redirects)


def test_normalize_columns_supports_units_aliases_and_spacing():
    raw = pd.DataFrame(columns=[
        ' Temperature (°C) ', 'Panel Age', 'Cleanliness Index',
        'Module Temp', 'Cloud Cover', 'Wind', 'Maintenance Counts',
    ])
    normalized = normalize_columns(raw)
    assert list(normalized.columns) == [
        'temperature', 'panel_age', 'soiling_ratio', 'module_temperature',
        'cloud_coverage', 'wind_speed', 'maintenance_count',
    ]


def test_upload_preprocessing_maps_unknown_categories_and_reports_them():
    row = _valid_row(
        string_id='STR-001', error_code='unexpected',
        installation_type='Rooftop',
    )
    encoded, notices = clean_for_upload(pd.DataFrame([row]))
    assert encoded.iloc[0]['string_id_B2'] == 0
    assert encoded.iloc[0]['string_id_C3'] == 0
    assert encoded.iloc[0]['string_id_D4'] == 0
    assert encoded.iloc[0]['error_code_missing'] == 1
    assert encoded.iloc[0]['installation_type_missing'] == 1
    assert [(notice['column'], notice['count']) for notice in notices] == [
        ('string_id', 1), ('error_code', 1), ('installation_type', 1)
    ]


def test_model_and_external_string_id_formats_are_supported():
    assert all(
        is_supported_string_identifier(value)
        for value in ('A1', 'B2', 'C3', 'D4')
    )
    assert all(
        is_supported_string_identifier(f'STR-{number:03}')
        for number in range(1, 101)
    )
    assert not is_supported_string_identifier('STR-000')
    assert not is_supported_string_identifier('STR-101')
    assert not is_supported_string_identifier('NOT-A-PANEL')

    rows = [
        _valid_row(string_id=value)
        for value in ('A1', 'B2', 'C3', 'D4', 'STR-001', 'STR-100')
    ]
    encoded, notices = clean_for_upload(pd.DataFrame(rows))
    assert list(encoded['string_id_B2']) == [0, 1, 0, 0, 0, 0]
    assert list(encoded['string_id_C3']) == [0, 0, 1, 0, 0, 0]
    assert list(encoded['string_id_D4']) == [0, 0, 0, 1, 0, 0]
    string_notice = next(
        notice for notice in notices if notice['column'] == 'string_id'
    )
    assert string_notice['count'] == 2
    assert string_notice['external_identifier_rows'] == 2


def test_genuinely_invalid_string_ids_use_explicit_baseline_fallback():
    row = _valid_row(string_id='STR-101')
    encoded, notices = clean_for_upload(pd.DataFrame([row]))
    assert not is_supported_string_identifier(row['string_id'])
    assert encoded.loc[0, [
        'string_id_B2', 'string_id_C3', 'string_id_D4'
    ]].tolist() == [0, 0, 0]
    notice = next(
        notice for notice in notices if notice['column'] == 'string_id'
    )
    assert notice['count'] == 1
    assert notice['external_identifier_rows'] == 0
    assert notice['default'] == 'A1'


def test_strict_inference_still_reports_allowed_category_values():
    with pytest.raises(ValueError, match='Allowed values: dual-axis, fixed'):
        clean_for_inference(pd.DataFrame([
            _valid_row(installation_type='Rooftop')
        ]))


def test_upload_with_only_numeric_features_uses_optional_defaults(app, client):
    _authenticate_client(app, client)
    row = _valid_row()
    for category in CATEGORICAL_LEVELS:
        row.pop(category)
    csv_data = pd.DataFrame([row]).to_csv(index=False).encode()
    response = _post_csv(client, csv_data, follow_redirects=True)
    assert response.status_code == 200
    assert b'Optional categorical column(s) absent' in response.data
    assert b'string_id (A1)' in response.data
    assert b'error_code (missing)' in response.data


def test_semicolon_bom_and_noncanonical_headers_upload(app, client):
    _authenticate_client(app, client)
    row = _valid_row()
    columns = {
        'Temperature (°C)': 'temperature',
        'Irradiance': 'irradiance',
        'Humidity': 'humidity',
        'Panel Age': 'panel_age',
        'Maintenance Count': 'maintenance_count',
        'Cleanliness Index': 'soiling_ratio',
        'Voltage': 'voltage',
        'Current': 'current',
        'Module Temp': 'module_temperature',
        'Cloud Cover': 'cloud_coverage',
        'Wind': 'wind_speed',
        'Pressure': 'pressure',
    }
    raw = pd.DataFrame([{header: row[key] for header, key in columns.items()}])
    csv_data = b'\xef\xbb\xbf' + raw.to_csv(index=False, sep=';').encode('utf-8')
    response = _post_csv(client, csv_data)
    assert response.status_code == 302
    assert '/upload/results/' in response.headers['Location']


def test_tab_delimited_upload_is_sniffed(app, client):
    _authenticate_client(app, client)
    csv_data = pd.DataFrame([_valid_row()]).to_csv(
        index=False, sep='\t'
    ).encode()
    response = _post_csv(client, csv_data)
    assert response.status_code == 302
    assert '/upload/results/' in response.headers['Location']


def test_latin1_encoded_csv_upload_uses_fallback(app, client):
    _authenticate_client(app, client)
    row = _valid_row()
    row['panel_id'] = 'panneau-é-1'
    csv_data = pd.DataFrame([row]).to_csv(index=False).encode('latin-1')
    response = _post_csv(client, csv_data)
    assert response.status_code == 302
    assert '/upload/results/' in response.headers['Location']


def test_csv_without_efficiency_uploads_and_downloads_predictions(app, client):
    _authenticate_client(app, client)
    token = _csrf(client, '/upload')
    csv_data = pd.DataFrame([_valid_row()]).to_csv(index=False).encode()
    response = client.post('/upload', data={
        'csrf_token': token,
        'file': (io.BytesIO(csv_data), 'panels.csv'),
    }, content_type='multipart/form-data')
    assert response.status_code == 302
    assert '/upload/results/' in response.headers['Location']
    result_page = client.get(response.headers['Location'])
    assert result_page.status_code == 200
    with app.app_context():
        dataset = Dataset.query.one()
        assert dataset.rows_count == 1
        assert dataset.mean_eff is not None
        download_url = f'/upload/{dataset.id}/download'
    download = client.get(download_url)
    assert download.status_code == 200
    output = pd.read_csv(io.BytesIO(download.data))
    assert {'predicted_efficiency', 'recommendation'} <= set(output.columns)
    assert 'efficiency' not in output.columns


def test_csv_missing_required_columns_shows_clear_error(app, client):
    _authenticate_client(app, client)
    token = _csrf(client, '/upload')
    csv_data = b'temperature\n25\n'
    response = client.post('/upload', data={
        'csrf_token': token,
        'file': (io.BytesIO(csv_data), 'incomplete.csv'),
    }, content_type='multipart/form-data', follow_redirects=True)
    assert response.status_code == 200
    assert b'Missing required column(s)' in response.data
    assert b'pressure' in response.data
    assert b'Columns found in your file: temperature' in response.data


def test_upload_accepts_case_and_whitespace_normalized_categories(app, client):
    _authenticate_client(app, client)
    row = _valid_row(string_id=' A1 ', error_code='e00',
                     installation_type='Fixed')
    response = _post_csv(
        client, pd.DataFrame([row]).to_csv(index=False).encode(),
        follow_redirects=True,
    )
    assert response.status_code == 200
    assert b'values the model does not know' not in response.data


def test_upload_unknown_category_is_defaulted_with_short_notice(app, client):
    _authenticate_client(app, client)
    row = _valid_row(installation_type='xyz')
    response = _post_csv(
        client, pd.DataFrame([row]).to_csv(index=False).encode(),
        follow_redirects=True,
    )
    assert response.status_code == 200
    assert b'installation_type: 1 rows' in response.data
    assert b'default group missing' in response.data
    notice = re.search(
        rb'<div class="message info">(.*?)</div>', response.data, re.DOTALL
    )
    assert notice is not None
    assert len(notice.group(1)) < 300


def test_upload_known_categories_do_not_produce_unknown_notice(app, client):
    _authenticate_client(app, client)
    rows = [
        _valid_row(string_id=string_id, error_code='E00',
                   installation_type='fixed')
        for string_id in ('A1', 'B2', 'C3', 'D4')
    ]
    response = _post_csv(
        client, pd.DataFrame(rows).to_csv(index=False).encode(),
        follow_redirects=True,
    )
    assert response.status_code == 200
    assert b'values the model does not know' not in response.data


def test_arbitrary_string_ids_are_kept_in_predictions_csv(app, client):
    _authenticate_client(app, client)
    ids = [f'STR-{index:03}' for index in range(1, 101)]
    rows = [_valid_row(string_id=value) for value in ids]
    response = _post_csv(
        client, pd.DataFrame(rows).to_csv(index=False).encode(),
        follow_redirects=True,
    )
    assert response.status_code == 200
    notices = re.findall(
        rb'<div class="message info">(.*?)</div>', response.data, re.DOTALL
    )
    unknown_notice = next(
        notice for notice in notices if b'string_id: 100 rows' in notice
    )
    assert b'STR-001' in unknown_notice
    assert b'and 95 more' in unknown_notice
    assert len(unknown_notice) < 300

    with app.app_context():
        dataset = Dataset.query.one()
        dataset_id = dataset.id
    download = client.get(f'/upload/{dataset_id}/download')
    output = pd.read_csv(io.BytesIO(download.data), dtype={'panel_id': str})
    assert download.status_code == 200
    assert len(output) == 100
    assert output['panel_id'].tolist() == ids


def test_upload_unknown_rooftop_installation_is_defaulted_with_notice(app, client):
    _authenticate_client(app, client)
    row = _valid_row(installation_type='Rooftop')
    response = _post_csv(
        client, pd.DataFrame([row]).to_csv(index=False).encode(),
        follow_redirects=True,
    )
    assert response.status_code == 200
    assert b'installation_type: 1 rows' in response.data
    assert b'Rooftop' in response.data
    assert b'default group missing' in response.data


def test_upload_template_is_login_protected_and_downloadable(app, client):
    anonymous = client.get('/upload/template')
    assert anonymous.status_code == 302
    assert anonymous.headers['Location'].endswith('/login?next=%2Fupload%2Ftemplate')

    _authenticate_client(app, client)
    response = client.get('/upload/template')
    assert response.status_code == 200
    assert response.headers['Content-Disposition'].endswith(
        'filename=sample_upload.csv'
    )
    sample = pd.read_csv(io.BytesIO(response.data))
    assert list(sample.columns) == [
        *NUMERIC_FEATURES, 'string_id', 'error_code', 'installation_type'
    ]
    assert len(sample) == 5


def test_training_dataset_csv_uploads_successfully(app, client):
    _authenticate_client(app, client)
    root = Path(__file__).resolve().parents[1]
    csv_data = (root / 'data' / 'train.csv').read_bytes()
    response = _post_csv(
        client, csv_data, filename='train.csv', follow_redirects=False
    )
    assert response.status_code == 302
    assert '/upload/results/' in response.headers['Location']


def test_single_panel_form_rejects_invalid_string_id(app, client):
    _authenticate_client(app, client)
    row = _valid_row(string_id='STR-001')
    row.update({'csrf_token': _csrf(client, '/predict'), 'mode': 'advanced'})
    response = client.post('/predict', data=row)
    assert response.status_code == 200
    assert b'String ID: choose a valid option.' in response.data


def test_maintenance_rule_uses_cleanliness_index():
    dusty = get_maintenance_recommendation(0.45, 0.5, 10)
    clean = get_maintenance_recommendation(0.45, 0.9, 10)
    assert dusty == (
        'Schedule panel cleaning — heavy soiling detected', 'warning', 'cleaning'
    )
    assert clean[0] != 'Schedule panel cleaning — heavy soiling detected'


def test_password_policy_rejects_weak_and_accepts_strong_password(client):
    weak = _register(client, 'weak-user', 'abcdef')
    assert b'Password must be at least 8 characters.' in weak.data
    strong = _register(client, 'strong-user', 'Solar@2026x')
    assert b'Registration successful' in strong.data


def test_registration_rejects_duplicate_email_case_insensitively(client):
    first = _register(
        client, 'email-one', email='Same@Example.test'
    )
    assert b'Registration successful' in first.data
    duplicate = _register(
        client, 'email-two', email='same@example.test'
    )
    assert b'That email address is already registered.' in duplicate.data


def test_registration_rejects_bad_email_format(client):
    response = _register(client, 'bad-email', email='not-an-email')
    assert b'Enter a valid email address.' in response.data


@pytest.mark.parametrize(
    ('identifier_field', 'identifier'),
    [('identifier', 'login-email@example.test'),
     ('identifier', 'login-username')],
)
def test_login_accepts_email_and_username(client, identifier_field, identifier):
    _register(client, 'login-username', email='login-email@example.test')
    token = _csrf(client, '/login')
    response = client.post('/login', data={
        'csrf_token': token,
        identifier_field: identifier,
        'password': 'Solar@2026x',
    })
    assert response.status_code == 302
    assert response.headers['Location'].endswith('/dashboard')


def test_maintenance_count_requires_whole_number(app, client):
    _authenticate_client(app, client)
    token = _csrf(client, '/predict')
    values = _valid_row(maintenance_count=4)
    values.update({'csrf_token': token, 'mode': 'advanced',
                   'maintenance_count': '4.5'})
    response = client.post('/predict', data=values)
    assert response.status_code == 200
    assert b'Maintenance count: enter a whole number.' in response.data
    with app.app_context():
        assert PredictionResult.query.count() == 0


def test_post_without_csrf_token_is_rejected(client):
    response = client.post('/register', data={
        'username': 'csrf-user',
        'password': 'Solar@2026x',
        'password_confirm': 'Solar@2026x',
    })
    assert response.status_code == 400


def test_secret_key_and_secure_cookie_configuration(monkeypatch):
    monkeypatch.delenv('SECRET_KEY', raising=False)
    monkeypatch.delenv('COOKIE_SECURE', raising=False)
    development = Flask('development-config-test')
    development.config['ENVIRONMENT'] = 'development'
    DevelopmentConfig.init_app(development)
    assert len(development.config['SECRET_KEY']) == 64
    assert development.config['SESSION_COOKIE_SECURE'] is False

    production = Flask('production-config-test')
    production.config['ENVIRONMENT'] = 'production'
    with pytest.raises(RuntimeError, match='non-default SECRET_KEY'):
        ProductionConfig.init_app(production)
    monkeypatch.setenv('SECRET_KEY', 'replace-with-a-generated-secret')
    with pytest.raises(RuntimeError, match='non-default SECRET_KEY'):
        ProductionConfig.init_app(production)

    monkeypatch.setenv('SECRET_KEY', 'generated-production-secret')
    ProductionConfig.init_app(production)
    assert production.config['SESSION_COOKIE_SECURE'] is True
    assert production.config['REMEMBER_COOKIE_SECURE'] is True


def test_production_app_rejects_missing_secret(monkeypatch):
    monkeypatch.delenv('SECRET_KEY', raising=False)
    with pytest.raises(RuntimeError, match='non-default SECRET_KEY'):
        create_app('production')


def test_health_is_json_and_does_not_query_or_write_database(app, client):
    statements = []

    def record_sql(_connection, _cursor, statement, *_args):
        statements.append(statement)

    with app.app_context():
        event.listen(db.engine, 'before_cursor_execute', record_sql)
        try:
            response = client.get('/health')
        finally:
            event.remove(db.engine, 'before_cursor_execute', record_sql)
    assert response.status_code == 200
    assert response.get_json() == {'status': 'ok'}
    assert statements == []


@pytest.mark.parametrize(
    ('trust_proxy', 'expected_ip'),
    [(False, '127.0.0.1'), (True, '198.51.100.23')],
)
def test_proxy_fix_controls_forwarded_ip_for_audit(
    monkeypatch, trust_proxy, expected_ip
):
    monkeypatch.setenv('TRUST_PROXY', str(trust_proxy).lower())
    app = create_app('testing')
    client = app.test_client()
    user_id = _create_user(app, f'proxy-{trust_proxy}')
    token = _csrf(client, '/login')
    response = client.post(
        '/login',
        data={
            'csrf_token': token,
            'identifier': f'proxy-{trust_proxy}',
            'password': 'Solar@2026x',
        },
        headers={'X-Forwarded-For': '198.51.100.23'},
    )
    assert response.status_code == 302
    with app.app_context():
        audit = AuditLog.query.filter_by(
            user_id=user_id, action='login_success'
        ).one()
        assert audit.ip == expected_ip


def test_startup_paths_and_model_work_from_an_unrelated_working_directory(
    tmp_path, monkeypatch
):
    monkeypatch.chdir(tmp_path)
    application = create_app('testing')
    assert Path(application.config['UPLOAD_FOLDER']).is_absolute()
    assert Path(application.instance_path).is_absolute()
    assert 0 <= predict_single(_valid_row()) <= 1


def test_create_all_startup_preserves_existing_database_data(
    tmp_path, monkeypatch
):
    database_path = (tmp_path / 'persisted.sqlite').as_posix()
    monkeypatch.setattr(
        TestingConfig, 'SQLALCHEMY_DATABASE_URI',
        f'sqlite:///{database_path}',
    )
    first = create_app('testing')
    with first.app_context():
        user = User(username='existing', email='existing@example.test')
        user.set_password('Solar@2026x')
        db.session.add(user)
        db.session.commit()
        user_id = user.id

    second = create_app('testing')
    with second.app_context():
        assert db.session.get(User, user_id).username == 'existing'


def test_deployment_manifests_and_ignore_rules_are_present():
    root = Path(__file__).resolve().parents[1]
    assert (root / '.python-version').read_text().strip() == '3.11'
    procfile = (root / 'Procfile').read_text()
    assert (
        'gunicorn --workers 1 --threads 2 --timeout 120 '
        '--bind 0.0.0.0:$PORT '
        '"app.app:create_app(\'production\')"'
    ) in procfile
    dockerfile = (root / 'Dockerfile').read_text()
    assert (
        'CMD ["gunicorn", "--workers", "1", "--threads", "2", '
        '"--timeout", "120", "--bind", "0.0.0.0:8000", '
        '"app.app:create_app(\'production\')"]'
    ) in dockerfile
    ignored = (root / '.gitignore').read_text()
    assert '*.db' in ignored


def test_prediction_ui_is_labeled_responsive_and_disables_on_submit(app, client):
    _authenticate_client(app, client)
    with client.session_transaction() as session:
        session['lang'] = 'kn'
    response = client.get('/predict')
    html = response.data.decode('utf-8')
    assert response.status_code == 200
    assert 'data-predicting="ಅಂದಾಜು ಮಾಡಲಾಗುತ್ತಿದೆ…"' in html
    assert "predictionButton.disabled = true" in html
    assert "predictionButtonLabel.textContent = predictionButton.dataset.predicting" in html
    for input_tag in re.findall(r'<input\b[^>]*>', html):
        if re.search(r'\btype="hidden"', input_tag):
            continue
        input_id = re.search(r'\bid="([^"]+)"', input_tag)
        assert input_id is not None, input_tag
        assert re.search(
            rf'<label\b[^>]*\bfor="{re.escape(input_id.group(1))}"',
            html,
        ), input_tag
    assert '@media (max-width:600px)' in html
    assert 'min-height: 44px' in html
    assert ':focus-visible' in html


def test_dashboard_and_history_have_translated_empty_states(app, client):
    _authenticate_client(app, client, 'empty-state-owner')
    with client.session_transaction() as session:
        session['lang'] = 'hi'
    dashboard = client.get('/dashboard')
    history = client.get('/history')
    assert 'अभी कोई अनुमान नहीं है'.encode() in dashboard.data
    assert 'अभी कोई अनुमान नहीं है'.encode() in history.data


@pytest.mark.parametrize(
    ('efficiency', 'soiling', 'expected_status', 'expected_class'),
    [
        (0.7, 0.9, 'Good', 'good'),
        (0.45, 0.5, 'Watch', 'watch'),
        (0.35, 0.9, 'Urgent', 'urgent'),
    ],
)
def test_results_status_matches_maintenance_recommendation_key(
    app, client, efficiency, soiling, expected_status, expected_class
):
    user_id = _authenticate_client(
        app, client, f'status-{expected_class}'
    )
    result_id = _create_prediction(
        app, user_id, efficiency=efficiency, soiling_ratio=soiling
    )
    page = client.get(f'/results/{result_id}')
    assert f'class="status-badge {expected_class}"'.encode() in page.data
    assert expected_status.encode() in page.data
    assert b'forecast-chart-summary' in page.data
    assert b'importance-chart-summary' in page.data
    assert b'chart_data_summary' not in page.data


def test_history_table_uses_a_keyboard_scroll_region(app, client):
    user_id = _authenticate_client(app, client, 'scroll-owner')
    _create_prediction(app, user_id)
    page = client.get('/history')
    assert b'class="table-scroll" tabindex="0" role="region"' in page.data
    assert b'min-width:680px' in page.data


def test_external_next_url_does_not_redirect_off_site(app, client):
    _create_user(app, 'redirect-user')
    token = _csrf(client, '/login?next=//evil.com')
    response = client.post('/login?next=//evil.com', data={
        'csrf_token': token,
        'username': 'redirect-user',
        'password': 'Solar@2026x',
        'next': '//evil.com',
    })
    assert response.status_code == 302
    assert not response.headers['Location'].startswith('//')
    assert 'evil.com' not in response.headers['Location']


def test_valid_out_of_training_range_input_gets_nonblocking_warning(app, client):
    _authenticate_client(app, client)
    token = _csrf(client, '/predict')
    values = _valid_row(irradiance=1600)
    values.update({'csrf_token': token, 'mode': 'advanced'})
    response = client.post('/predict', data=values)
    assert response.status_code == 302
    result_page = client.get(response.headers['Location'])
    assert b'Input outside training range' in result_page.data
    assert b'Prediction result #' in result_page.data


def test_user_cannot_open_another_users_prediction(app, client):
    with app.app_context():
        owner = User(username='result-owner', email='result-owner@example.test')
        owner.set_password('Solar@2026x')
        other = User(username='other-user', email='other-user@example.test')
        other.set_password('Solar@2026x')
        db.session.add_all([owner, other])
        db.session.flush()
        result = PredictionResult(
            user_id=owner.id,
            input_data='{"soiling_ratio":0.8,"panel_age":10}',
            predicted_efficiency=0.5,
            maintenance_recommendation='test',
        )
        db.session.add(result)
        db.session.commit()
        result_id = result.id
        other_id = other.id
    with client.session_transaction() as session:
        session['_user_id'] = str(other_id)
        session['_fresh'] = True
    assert client.get(f'/results/{result_id}').status_code == 404


def test_simple_and_advanced_prediction_reports_and_logout(app, client):
    _register(client, 'end-to-end-user')
    login_token = _csrf(client, '/login')
    login_response = client.post('/login', data={
        'csrf_token': login_token,
        'username': 'end-to-end-user',
        'password': 'Solar@2026x',
    })
    assert login_response.status_code == 302
    assert b'Simple mode' in client.get('/predict').data

    simple_token = _csrf(client, '/predict')
    simple_response = client.post('/predict', data={
        'csrf_token': simple_token,
        'mode': 'simple',
        'panel_age': '10',
        'installation_type': 'fixed',
        'cleanliness': '0.9',
        'maintenance_count': '4',
    })
    assert simple_response.status_code == 302
    simple_result_url = simple_response.headers['Location']
    simple_page = client.get(simple_result_url)
    assert b'Prediction mode: <strong>Simple</strong>' in simple_page.data
    assert b'degradationChart' in simple_page.data
    assert b'importanceChart' in simple_page.data

    result_id = simple_result_url.rsplit('/', 1)[1]
    pdf = client.get(f'/reports/{result_id}/pdf')
    excel = client.get(f'/reports/{result_id}/excel')
    assert pdf.status_code == 200 and pdf.data.startswith(b'%PDF')
    assert excel.status_code == 200
    workbook = openpyxl.load_workbook(io.BytesIO(excel.data), data_only=True)
    summary = list(workbook['Summary'].values)
    assert any('Efficiency change per year' in str(row[0]) for row in summary)
    assert any('Years until efficiency falls below 0.40' in str(row[0]) for row in summary)

    advanced_token = _csrf(client, '/predict')
    advanced_values = _valid_row()
    advanced_values.update({'csrf_token': advanced_token, 'mode': 'advanced'})
    advanced_response = client.post('/predict', data=advanced_values)
    assert advanced_response.status_code == 302
    advanced_page = client.get(advanced_response.headers['Location'])
    assert b'Prediction mode: <strong>Advanced</strong>' in advanced_page.data

    logout_token = _csrf(client, '/dashboard')
    logout = client.post('/logout', data={'csrf_token': logout_token})
    assert logout.status_code == 302
    assert logout.headers['Location'].endswith('/login')
    with app.app_context():
        actions = {entry.action for entry in AuditLog.query.all()}
    assert {'register', 'login_success', 'prediction', 'report_download', 'logout'} <= actions


def test_translation_helper_falls_back_for_missing_key_and_language(app):
    with app.app_context():
        assert t('not_a_translation_key', lang='kn') == 'not_a_translation_key'
        assert t('login', lang='xx') == TRANSLATIONS['login']['en']


def test_template_and_report_translation_keys_cover_all_supported_languages():
    from pathlib import Path

    root = Path(__file__).resolve().parents[1]
    files = list((root / 'app' / 'templates').glob('*.html'))
    files += list((root / 'app' / 'routes').glob('*.py'))
    files.append(root / 'app' / 'utils' / 'report_generator.py')
    referenced = set()
    for path in files:
        referenced.update(re.findall(
            r"""\bt\(\s*['"]([^'"]+)['"]""",
            path.read_text(encoding='utf-8'),
        ))
    assert referenced <= set(TRANSLATIONS)
    assert all(
        set(SUPPORTED_LANGS) <= set(TRANSLATIONS[key])
        for key in referenced
    )


def test_language_switch_allowlist_and_csrf(app, client):
    _authenticate_client(app, client)
    token = _csrf(client, '/dashboard')
    rejected = client.post('/set-language/xx', data={
        'csrf_token': token,
        'next': '/dashboard',
    })
    assert rejected.status_code == 302
    with client.session_transaction() as session:
        assert session.get('lang', 'en') == 'en'

    missing_token = client.post('/set-language/hi', data={'next': '/dashboard'})
    assert missing_token.status_code == 400
    with client.session_transaction() as session:
        assert session.get('lang', 'en') == 'en'


def test_saved_language_persists_across_logout_and_login(app, client):
    _register(client, 'language-user')
    login_token = _csrf(client, '/login')
    client.post('/login', data={
        'csrf_token': login_token,
        'identifier': 'language-user',
        'password': 'Solar@2026x',
    })
    switch_token = _csrf(client, '/dashboard')
    client.post('/set-language/te', data={
        'csrf_token': switch_token,
        'next': '/dashboard',
    })
    with app.app_context():
        user = User.query.filter_by(username='language-user').one()
        assert user.language == 'te'

    logout_token = _csrf(client, '/dashboard')
    client.post('/logout', data={'csrf_token': logout_token})
    relogin_token = _csrf(client, '/login')
    client.post('/login', data={
        'csrf_token': relogin_token,
        'identifier': 'language-user',
        'password': 'Solar@2026x',
    })
    with client.session_transaction() as session:
        assert session['lang'] == 'te'
    assert b'lang="te"' in client.get('/dashboard').data


@pytest.mark.parametrize('ui_lang', list(SUPPORTED_LANGS))
def test_client_workflow_downloads_report_in_language_other_than_ui(
    app, client, ui_lang
):
    _register(client, f'workflow-{ui_lang}')
    login_token = _csrf(client, '/login')
    client.post('/login', data={
        'csrf_token': login_token,
        'identifier': f'workflow-{ui_lang}',
        'password': 'Solar@2026x',
    })
    prediction_token = _csrf(client, '/predict')
    prediction = client.post('/predict', data={
        'csrf_token': prediction_token,
        'mode': 'simple',
        'panel_age': '10',
        'installation_type': 'fixed',
        'cleanliness': '0.9',
        'maintenance_count': '4',
    })
    assert prediction.status_code == 302
    results_url = prediction.headers['Location']
    result_page = client.get(results_url)
    assert b'lang="en"' in result_page.data

    switch_token = _csrf(client, results_url)
    client.post(f'/set-language/{ui_lang}', data={
        'csrf_token': switch_token,
        'next': results_url,
    })
    localized_page = client.get(results_url)
    assert f'lang="{ui_lang}"'.encode() in localized_page.data

    report_lang = {
        'en': 'kn', 'kn': 'hi', 'hi': 'te', 'te': 'ta', 'ta': 'en',
    }[ui_lang]
    result_id = results_url.rsplit('/', 1)[1]
    pdf = client.get(f'/reports/{result_id}/pdf?lang={report_lang}')
    excel = client.get(f'/reports/{result_id}/excel?lang={report_lang}')
    assert pdf.status_code == 200 and b'%PDF' in pdf.data[:10]
    assert f'NotoSans'.encode() in pdf.data
    workbook = openpyxl.load_workbook(io.BytesIO(excel.data), data_only=True)
    assert workbook.sheetnames[0] == t('sheet_summary', lang=report_lang)


def test_recommendation_is_translated_but_database_text_stays_english(app, client):
    user_id = _authenticate_client(app, client, 'recommendation-owner')
    result_id = _create_prediction(app, user_id)
    with client.session_transaction() as session:
        session['lang'] = 'ta'
    page = client.get(f'/results/{result_id}')
    assert 'பலகையைச் சுத்தம் செய்ய திட்டமிடுங்கள்'.encode() in page.data
    with app.app_context():
        saved = db.session.get(PredictionResult, result_id)
        assert saved.maintenance_recommendation == (
            'Schedule panel cleaning — heavy soiling detected'
        )


@pytest.mark.parametrize('lang', list(SUPPORTED_LANGS))
def test_localized_pdf_and_excel_reports_for_each_language(app, client, lang):
    user_id = _authenticate_client(app, client, f'report-{lang}')
    result_id = _create_prediction(app, user_id)
    with app.app_context():
        result = db.session.get(PredictionResult, result_id)
        pdf_bytes = generate_pdf(result, lang=lang)
        excel_bytes = generate_excel(result, lang=lang)

    assert pdf_bytes.startswith(b'%PDF')
    assert b'NotoSans' in pdf_bytes
    assert b'/FontFile2' in pdf_bytes
    workbook = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
    assert workbook.sheetnames == [
        t('sheet_summary', lang=lang),
        t('sheet_input_data', lang=lang),
    ]
    summary = list(workbook[workbook.sheetnames[0]].values)
    assert any(row[0] == t('recommendation', lang=lang) for row in summary)
    assert any(row[1] == t('cleaning', lang=lang) for row in summary)
    data = list(workbook[workbook.sheetnames[1]].values)
    assert data[1][0] == t('temperature', lang=lang)


def test_invalid_report_language_falls_back_to_current_language(app, client):
    user_id = _authenticate_client(app, client, 'fallback-report-owner')
    result_id = _create_prediction(app, user_id)
    with client.session_transaction() as session:
        session['lang'] = 'ta'
    response = client.get(f'/reports/{result_id}/pdf?lang=zz')
    assert response.status_code == 200
    assert b'NotoSansTamil' in response.data


@pytest.mark.parametrize('lang', list(SUPPORTED_LANGS))
@pytest.mark.parametrize('format_name', ['pdf', 'excel'])
def test_other_users_cannot_download_reports_in_any_language(
    app, client, lang, format_name
):
    owner_id = _create_user(app, f'private-report-owner-{lang}-{format_name}')
    other_id = _create_user(app, f'private-report-other-{lang}-{format_name}')
    result_id = _create_prediction(app, owner_id)
    with client.session_transaction() as session:
        session['_user_id'] = str(other_id)
        session['_fresh'] = True
    assert client.get(
        f'/reports/{result_id}/{format_name}?lang={lang}'
    ).status_code == 404


@pytest.mark.parametrize('lang', ['te', 'ta'])
def test_long_indic_recommendation_wraps_without_pdf_failure(app, lang, monkeypatch):
    user_id = _create_user(app, f'long-report-{lang}')
    result_id = _create_prediction(app, user_id)
    repeated = TRANSLATIONS['cleaning'][lang] + ' '
    monkeypatch.setitem(
        TRANSLATIONS['cleaning'], lang, repeated * 45
    )
    with app.app_context():
        result = db.session.get(PredictionResult, result_id)
        content = generate_pdf(result, lang=lang)
    assert content.startswith(b'%PDF')
