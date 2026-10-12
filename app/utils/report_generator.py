"""Generate localized PDF and Excel exports for a saved prediction."""

from datetime import datetime
from io import BytesIO
from pathlib import Path

import openpyxl
from fpdf import FPDF
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.utils.maintenance_rules import get_maintenance_recommendation
from app.utils.translations import SUPPORTED_LANGS, t

FONT_DIR = Path(__file__).resolve().parents[1] / 'static' / 'fonts'
FONT_FILES = {
    'en': ('NotoSans', 'NotoSans-Regular.ttf', 'NotoSans-Bold.ttf'),
    'kn': ('NotoSansKannada', 'NotoSansKannada-Regular.ttf', 'NotoSansKannada-Bold.ttf'),
    'hi': ('NotoSansDevanagari', 'NotoSansDevanagari-Regular.ttf', 'NotoSansDevanagari-Bold.ttf'),
    'te': ('NotoSansTelugu', 'NotoSansTelugu-Regular.ttf', 'NotoSansTelugu-Bold.ttf'),
    'ta': ('NotoSansTamil', 'NotoSansTamil-Regular.ttf', 'NotoSansTamil-Bold.ttf'),
}

SEVERITY_COLORS = {
    'critical': (198, 40, 40),
    'warning': (230, 81, 0),
    'ok': (46, 125, 50),
}
SEVERITY_HEX = {'critical': 'C62828', 'warning': 'E65100', 'ok': '2E7D32'}
CATEGORY_LABELS = {
    'string_id': {
        'A1': 'string_a1', 'B2': 'string_b2', 'C3': 'string_c3', 'D4': 'string_d4',
    },
    'error_code': {
        'E00': 'error_none', 'E01': 'error_e01', 'E02': 'error_e02',
        'missing': 'opt_missing',
    },
    'installation_type': {
        'fixed': 'opt_fixed', 'tracking': 'opt_tracking',
        'dual-axis': 'opt_dual_axis', 'missing': 'opt_missing',
    },
}


def _localized_inputs(input_data, lang):
    labels = {}
    values = {}
    for feature, value in input_data.items():
        labels[feature] = t(
            'soiling_label' if feature == 'soiling_ratio' else feature, lang=lang
        )
        if feature in CATEGORY_LABELS:
            value = t(CATEGORY_LABELS[feature].get(value, 'opt_missing'), lang=lang)
        values[feature] = value
    return labels, values


def _result_details(result, lang):
    input_data = result.get_input_data()
    prediction_mode = input_data.pop('prediction_mode', 'advanced')
    _, severity, recommendation_key = get_maintenance_recommendation(
        result.predicted_efficiency,
        input_data['soiling_ratio'],
        input_data['panel_age'],
    )
    from ml.degradation import forecast_degradation

    forecast = forecast_degradation(input_data)
    labels, values = _localized_inputs(input_data, lang)
    threshold = (
        t('not_within_range', lang=lang)
        if forecast['years_to_threshold'] is None
        else t(
            'years_value',
            lang=lang,
            value=f"{forecast['years_to_threshold']:.1f}",
        )
    )
    return {
        'input_data': input_data,
        'prediction_mode': prediction_mode,
        'severity': severity,
        'recommendation': t(recommendation_key, lang=lang),
        'forecast': forecast,
        'threshold_text': threshold,
        'labels': labels,
        'values': values,
    }


class _SolarPDF(FPDF):
    def __init__(self, lang, family):
        super().__init__()
        self.lang = lang
        self.family = family

    def header(self):
        self.set_fill_color(74, 91, 177)
        self.rect(0, 0, 210, 22, 'F')
        self.set_xy(10, 5)
        self.set_font(self.family, 'B', 12)
        self.set_text_color(255, 255, 255)
        self.multi_cell(
            190, 8, t('app_name', lang=self.lang),
            new_x='LMARGIN', new_y='NEXT',
        )
        self.set_text_color(0, 0, 0)
        self.set_y(27)
        self.set_x(self.l_margin)

    def footer(self):
        self.set_y(-13)
        self.set_font(self.family, '', 7)
        self.set_text_color(120, 120, 120)
        self.cell(
            0, 8,
            t('report_footer', lang=self.lang, page=self.page_no()),
            align='C',
        )


def _write_wrapped(pdf, height, text, **kwargs):
    pdf.set_x(pdf.l_margin)
    pdf.multi_cell(
        0, height, text, new_x='LMARGIN', new_y='NEXT', **kwargs
    )


def generate_pdf(result, lang='en') -> bytes:
    """Render a prediction report using the requested language's Noto font."""
    lang = lang if lang in SUPPORTED_LANGS else 'en'
    family, regular_font, bold_font = FONT_FILES[lang]
    pdf = _SolarPDF(lang, family)
    pdf.add_font(family, '', str(FONT_DIR / regular_font))
    pdf.add_font(family, 'B', str(FONT_DIR / bold_font))
    pdf.add_font('NotoLatin', '', str(FONT_DIR / 'NotoSans-Regular.ttf'))
    pdf.add_font('NotoLatin', 'B', str(FONT_DIR / 'NotoSans-Bold.ttf'))
    pdf.set_fallback_fonts(['NotoLatin'], exact_match=False)
    pdf.set_text_shaping(True)
    pdf.set_auto_page_break(auto=True, margin=17)
    pdf.alias_nb_pages()

    details = _result_details(result, lang)
    forecast = details['forecast']
    color = SEVERITY_COLORS[details['severity']]
    pdf.add_page()
    pdf.set_font(family, 'B', 15)
    _write_wrapped(pdf, 9, t('report_title', lang=lang))
    pdf.set_font(family, '', 9)
    _write_wrapped(pdf, 6, f"{t('result_id', lang=lang)}: #{result.id}")
    _write_wrapped(
        pdf, 6,
        f"{t('generated', lang=lang)}: "
        f"{result.created_at.strftime('%Y-%m-%d %H:%M:%S UTC')}",
    )
    _write_wrapped(
        pdf, 6,
        f"{t('report_prediction_mode', lang=lang)}: "
        f"{t(details['prediction_mode'], lang=lang)}",
    )
    pdf.ln(3)

    pdf.set_fill_color(245, 247, 255)
    pdf.set_text_color(*color)
    pdf.set_font(family, 'B', 12)
    _write_wrapped(
        pdf, 9,
        f"{t('predicted_efficiency', lang=lang)}: "
        f"{result.predicted_efficiency:.4f} "
        f"({result.predicted_efficiency * 100:.1f}%)",
        fill=True,
    )
    pdf.set_text_color(*color)
    pdf.set_font(family, 'B', 10)
    severity_label = (
        t('severity_ok', lang=lang)
        if details['severity'] == 'ok'
        else t(details['severity'], lang=lang)
    )
    _write_wrapped(pdf, 7, f"{severity_label}: {details['recommendation']}")
    pdf.set_text_color(0, 0, 0)
    pdf.ln(2)

    pdf.set_font(family, 'B', 11)
    _write_wrapped(pdf, 7, t('degradation_forecast', lang=lang))
    pdf.set_font(family, '', 9)
    _write_wrapped(
        pdf, 6,
        f"{t('loss_card', lang=lang)}: {forecast['loss_per_year']:+.5f}",
    )
    _write_wrapped(
        pdf, 6,
        f"{t('threshold_card', lang=lang)}: {details['threshold_text']}",
    )
    pdf.set_font(family, '', 8)
    _write_wrapped(pdf, 5, t('forecast_explanation', lang=lang))
    pdf.ln(2)

    pdf.set_font(family, 'B', 11)
    _write_wrapped(pdf, 7, t('input_values', lang=lang))
    pdf.set_font(family, '', 8)
    for feature, value in details['values'].items():
        shown = f'{value:.4f}' if isinstance(value, float) else str(value)
        _write_wrapped(pdf, 5, f"{details['labels'][feature]}: {shown}")

    pdf.ln(2)
    pdf.set_font(family, '', 7)
    _write_wrapped(pdf, 5, t('report_model_note', lang=lang))
    return bytes(pdf.output())


def _valid_sheet_name(value):
    for forbidden in (':', '\\', '/', '?', '*', '[', ']'):
        value = value.replace(forbidden, ' ')
    return value[:31] or 'Report'


def generate_excel(result, lang='en') -> bytes:
    """Build a translated two-sheet workbook with readable wrapped cell text."""
    lang = lang if lang in SUPPORTED_LANGS else 'en'
    details = _result_details(result, lang)
    forecast = details['forecast']
    severity = details['severity']
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = _valid_sheet_name(t('sheet_summary', lang=lang))
    border_side = Side(style='thin', color='CCCCCC')
    border = Border(
        left=border_side, right=border_side, top=border_side, bottom=border_side
    )
    title_fill = PatternFill('solid', fgColor='667EEA')
    title_font = Font(name='Noto Sans', bold=True, size=13, color='FFFFFF')
    body_font = Font(name='Noto Sans', size=10)

    ws.merge_cells('A1:B1')
    ws['A1'] = t('report_title', lang=lang)
    ws['A1'].font = title_font
    ws['A1'].fill = title_fill
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 30
    recommendation_severity = (
        t('severity_ok', lang=lang)
        if severity == 'ok'
        else t(severity, lang=lang)
    )
    rows = [
        (t('result_id', lang=lang), f'#{result.id}'),
        (t('generated', lang=lang), result.created_at.strftime('%Y-%m-%d %H:%M:%S UTC')),
        (
            t('report_prediction_mode', lang=lang),
            t(details['prediction_mode'], lang=lang),
        ),
        (
            t('predicted_efficiency', lang=lang),
            f'{result.predicted_efficiency:.4f} ({result.predicted_efficiency * 100:.1f}%)',
        ),
        (t('severity', lang=lang), recommendation_severity),
        (t('recommendation', lang=lang), details['recommendation']),
        (
            t('loss_card', lang=lang),
            f"{forecast['loss_per_year']:+.5f}",
        ),
        (t('threshold_card', lang=lang), details['threshold_text']),
    ]
    for row_number, (label, value) in enumerate(rows, start=2):
        label_cell = ws.cell(row=row_number, column=1, value=label)
        value_cell = ws.cell(row=row_number, column=2, value=value)
        label_cell.font = Font(name='Noto Sans', bold=True, size=10)
        value_cell.font = body_font
        label_cell.alignment = value_cell.alignment = Alignment(
            vertical='top', wrap_text=True
        )
        label_cell.border = value_cell.border = border
        if label in (t('severity', lang=lang), t('recommendation', lang=lang)):
            value_cell.font = Font(name='Noto Sans', bold=True, size=10,
                                   color=SEVERITY_HEX[severity])
        ws.row_dimensions[row_number].height = 32
    ws.column_dimensions['A'].width = 38
    ws.column_dimensions['B'].width = 72
    note_row = len(rows) + 3
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=2)
    ws.cell(row=note_row, column=1, value=t('report_model_note', lang=lang))
    ws.cell(row=note_row, column=1).font = Font(
        name='Noto Sans', italic=True, size=9, color='666666'
    )
    ws.cell(row=note_row, column=1).alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[note_row].height = 50

    data_sheet = wb.create_sheet(_valid_sheet_name(t('sheet_input_data', lang=lang)))
    data_sheet.append([t('report_input_feature', lang=lang), t('input_values', lang=lang)])
    for cell in data_sheet[1]:
        cell.font = Font(name='Noto Sans', bold=True, color='FFFFFF')
        cell.fill = title_fill
        cell.border = border
        cell.alignment = Alignment(wrap_text=True)
    for feature, value in details['values'].items():
        shown = f'{value:.4f}' if isinstance(value, float) else str(value)
        data_sheet.append([details['labels'][feature], shown])
        for cell in data_sheet[data_sheet.max_row]:
            cell.font = body_font
            cell.border = border
            cell.alignment = Alignment(wrap_text=True, vertical='top')
        data_sheet.row_dimensions[data_sheet.max_row].height = 30
    data_sheet.column_dimensions['A'].width = 36
    data_sheet.column_dimensions['B'].width = 48
    data_sheet.freeze_panes = 'A2'
    data_sheet.sheet_view.showGridLines = False

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
